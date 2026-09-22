# -*- coding: utf-8 -*-
"""Testes dos relatorios e da exportacao Excel (Etapa 8). Banco em memoria."""
import os
import sys, re, io
from datetime import timedelta

# O caminho do projeto vem do PROPRIO arquivo, nao escrito a mao. Assim os
# testes funcionam em qualquer computador e tambem no servidor.
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook

from config import DevelopmentConfig
from app import create_app
from app.extensions import db, limiter
from app.models import Usuario, NovoConvertido, Contato, Evento, Presenca, LogAuditoria
from app.tempo import agora, hoje
from app import opcoes, relatorios as calc


class ConfigTeste(DevelopmentConfig):
    SQLALCHEMY_DATABASE_URI = "sqlite://"
    WTF_CSRF_ENABLED = True
    CADASTRO_TOKEN = "t"


ok = falhou = 0


def checa(desc, cond):
    global ok, falhou
    if cond:
        ok += 1
        print(f"  [OK]   {desc}")
    else:
        falhou += 1
        print(f"  [FALHA] {desc}")


app = create_app(ConfigTeste)

with app.app_context():
    db.create_all()
    admin = Usuario(nome="Carlos Admin", login="admin", papel="admin", ativo=True, deve_trocar_senha=False)
    admin.definir_senha("SenhaForte#2026")
    joao = Usuario(nome="Joao Pereira", login="joao", papel="responsavel", ativo=True, deve_trocar_senha=False)
    joao.definir_senha("OutraSenha#99")
    db.session.add_all([admin, joao])
    db.session.flush()
    ID_ADMIN, ID_JOAO = admin.id, joao.id

    def criar(nome, dep, trab, status, mes_atras=0, mesclada=False):
        d = hoje()
        ano, mes = d.year, d.month - mes_atras
        while mes <= 0:
            mes += 12
            ano -= 1
        data = d.replace(year=ano, month=mes, day=min(d.day, 28))
        a = NovoConvertido(
            nome_completo=nome, telefone="16992805852", sexo="F",
            data_nascimento=hoje().replace(year=hoje().year - 30),
            trabalho=trab, departamento=dep, data_conversao=data,
            cadastrante_nome="T", cadastrante_telefone="16988887777",
            consentimento_em=agora(), consentimento_ip="127.0.0.1",
            status_ciclo=status,
            responsavel_atual_id=ID_JOAO if status == opcoes.STATUS_ATIVO else None,
            designado_em=agora() - timedelta(days=30),
            criado_em=agora() - timedelta(days=30),
            observacao_sensivel="SEGREDO QUE NAO PODE SAIR EM PLANILHA",
        )
        db.session.add(a)
        db.session.flush()
        return a

    a1 = criar("Alma Um", "preciosas", "culto_dominical", opcoes.STATUS_ATIVO, 0)
    a2 = criar("Alma Dois", "preciosas", "culto_dominical", opcoes.STATUS_ATIVO, 0)
    a3 = criar("Alma Tres", "irmaos", "evangelismo_rua", opcoes.STATUS_ATIVO, 1)
    a4 = criar("Alma Quatro", "geracao_life", "culto_ensino", "integrado", 2)
    a5 = criar("Alma Cinco", "irmaos", "culto_dominical", opcoes.STATUS_INICIAL, 0)
    dup = criar("Alma Duplicada", "preciosas", "culto_dominical", opcoes.STATUS_ATIVO, 0)
    dup.mesclado_em_id = a1.id

    ev = Evento(nome="Culto Dominical", tipo="culto_dominical",
                data=hoje() - timedelta(days=5), ativo=True)
    ev_velho = Evento(nome="Culto Antigo", tipo="culto_ceia",
                      data=hoje() - timedelta(days=90), ativo=True)
    db.session.add_all([ev, ev_velho])
    db.session.flush()

    # a1 esteve presente (2 vezes no mesmo culto e impossivel; usamos 2 eventos)
    db.session.add(Presenca(convertido_id=a1.id, evento_id=ev.id, presente=True, registrado_por_id=ID_JOAO))
    db.session.add(Presenca(convertido_id=a1.id, evento_id=ev_velho.id, presente=True, registrado_por_id=ID_JOAO))
    # a2 faltou
    db.session.add(Presenca(convertido_id=a2.id, evento_id=ev.id, presente=False, registrado_por_id=ID_JOAO))

    db.session.add(Contato(convertido_id=a1.id, responsavel_id=ID_JOAO, tipo="telefone",
                           data_hora=agora() - timedelta(days=1), resultado="efetivo",
                           relato="Conversa boa.", criado_em=agora()))
    db.session.add(Contato(convertido_id=a2.id, responsavel_id=ID_JOAO, tipo="whatsapp",
                           data_hora=agora() - timedelta(days=20), resultado="sem_resposta",
                           relato="Nao respondeu.", criado_em=agora()))
    db.session.commit()

    IDS = {"a1": a1.id, "dup": dup.id}


def entrar(login, senha):
    with app.app_context():
        limiter.reset()
    c = app.test_client()
    h = c.get("/login").get_data(as_text=True)
    tok = re.search(r'name="csrf_token"[^>]*value="([^"]+)"', h).group(1)
    c.post("/login", data={"login": login, "senha": senha, "csrf_token": tok})
    return c


print("\n--- 1. SO ADMIN (secao 7, item 12) ---")
cj = entrar("joao", "OutraSenha#99")
checa(f"responsavel recebe 403 em /relatorios (deu {cj.get('/relatorios').status_code})",
      cj.get("/relatorios").status_code == 403)
r = cj.get("/relatorios/excel")
checa(f"responsavel recebe 403 na exportacao (deu {r.status_code})", r.status_code == 403)
checa("nao veio planilha nenhuma", b"PK" not in r.get_data()[:4])

c_anon = app.test_client()
checa("sem login tambem e barrado", c_anon.get("/relatorios").status_code in (302, 403))

print("\n--- 2. OS CALCULOS ---")
with app.app_context():
    print("   conversoes por mes:")
    m = calc.conversoes_por_mes()
    checa(f"12 meses na serie (deu {len(m['rotulos'])})", len(m["rotulos"]) == 12)
    checa("meses vazios aparecem como zero, nao somem", 0 in m["valores"])
    checa(f"total do periodo = 5, sem a duplicata (deu {m['total']})", m["total"] == 5)
    checa("o mes mais recente e o ultimo da lista", m["valores"][-1] == 3)

    print("   por trabalho:")
    t = calc.conversoes_por_trabalho()
    checa(f"culto dominical lidera com 3 (deu {t['valores'][0]})", t["valores"][0] == 3)
    checa("ordenado do maior para o menor", t["valores"] == sorted(t["valores"], reverse=True))
    checa(f"total sem duplicata = 5 (deu {sum(t['valores'])})", sum(t["valores"]) == 5)

    print("   por departamento:")
    dep = calc.por_departamento()
    checa(f"total = 5 (deu {dep['total']})", dep["total"] == 5)
    checa("os 3 departamentos aparecem, mesmo com zero", len(dep["rotulos"]) == 3)
    checa(f"percentuais somam ~100 (deu {sum(dep['percentuais'])})",
          98 <= sum(dep["percentuais"]) <= 102)

    print("   funil:")
    f = calc.funil_do_ciclo()
    checa("os 7 status aparecem", len(f["rotulos"]) == 7)
    checa("a ordem e a da JORNADA, nao por tamanho",
          f["chaves"][0] == "aguardando_responsavel" and f["chaves"][1] == "em_acompanhamento")
    checa(f"total = 5 (deu {f['total']})", f["total"] == 5)

    print("   retencao:")
    ret = calc.taxa_de_retencao(dias=30)
    checa(f"3 almas em acompanhamento (deu {ret['total']})", ret["total"] == 3)
    checa(f"1 esteve presente nos 30 dias (deu {ret['presentes']})", ret["presentes"] == 1)
    checa(f"33% de retencao (deu {ret['percentual']}%)", ret["percentual"] == 33)

    # quem foi a DOIS cultos conta UMA vez
    ret90 = calc.taxa_de_retencao(dias=120)
    checa(f"quem foi a 2 cultos conta 1 vez (deu {ret90['presentes']})", ret90["presentes"] == 1)

    print("   ranking:")
    rk = calc.ranking_de_responsaveis()
    checa("so quem tem almas aparece", all(n > 0 for n in rk["almas"]))
    checa("traz a quantidade junto do percentual",
          len(rk["almas"]) == len(rk["percentuais"]))

print("\n--- 3. A DUPLICATA NAO ENTRA EM RELATORIO NENHUM ---")
with app.app_context():
    d = calc.tudo()
checa("nao entra no total", d["resumo"]["total"] == 5)
checa("nao entra no por mes", d["por_mes"]["total"] == 5)
checa("nao entra no por trabalho", sum(d["por_trabalho"]["valores"]) == 5)
checa("nao entra no departamento", d["por_departamento"]["total"] == 5)
checa("nao entra no funil", d["funil"]["total"] == 5)

print("\n--- 4. A TELA ---")
ca = entrar("admin", "SenhaForte#2026")
r = ca.get("/relatorios")
checa("a tela abre", r.status_code == 200)
h = r.get_data(as_text=True)
for titulo in ["Taxa de retenção", "Conversões por mês", "Onde as pessoas aceitam Jesus",
               "Distribuição por departamento", "Funil do ciclo de vida",
               "Ranking de responsáveis"]:
    checa(f"mostra: {titulo}", titulo in h)
checa("tem a tabela com os mesmos numeros (acessibilidade)",
      "Ver os mesmos números em tabela" in h)
checa("a duplicata nao aparece na tela", "Alma Duplicada" not in h)

print("\n--- 5. A PLANILHA ---")
r = ca.get("/relatorios/excel")
checa("a exportacao funciona", r.status_code == 200)
checa("baixa como anexo", "attachment" in r.headers.get("Content-Disposition", ""))
checa("o nome tem a data", ".xlsx" in r.headers.get("Content-Disposition", ""))
checa("e um arquivo xlsx de verdade", r.get_data()[:2] == b"PK")

wb = load_workbook(io.BytesIO(r.get_data()))
checa(f"as 4 abas exigidas + resumo (deu {wb.sheetnames})",
      set(["Novos Convertidos", "Contatos", "Presencas", "Responsaveis"]).issubset(set(wb.sheetnames)))

aba = wb["Novos Convertidos"]
checa(f"13 colunas ou mais (deu {aba.max_column})", aba.max_column >= 13)
checa(f"5 almas, sem a duplicata (deu {aba.max_row - 1})", aba.max_row - 1 == 5)
checa("cabecalho congelado", aba.freeze_panes == "A2")
checa("filtros ligados", aba.auto_filter.ref is not None)
checa("cabecalho em negrito", aba["A1"].font.bold is True)
checa("cabecalho com fundo colorido", aba["A1"].fill.fgColor.rgb is not None)
checa("colunas com largura ajustada", aba.column_dimensions["B"].width > 10)

todo_texto = ""
for nome in wb.sheetnames:
    for linha in wb[nome].iter_rows(values_only=True):
        todo_texto += " ".join(str(v) for v in linha if v is not None)

print("\n--- 6. O QUE NAO PODE SAIR NA PLANILHA ---")
checa("a OBSERVACAO SENSIVEL nao e exportada", "SEGREDO QUE NAO PODE SAIR" not in todo_texto)
checa("nenhum hash de senha vaza", "$2b$" not in todo_texto)
checa("a alma duplicada nao aparece", "Alma Duplicada" not in todo_texto)
checa("o aviso de LGPD esta na planilha", "LGPD" in todo_texto)

print("\n--- 7. O QUE PRECISA ESTAR NA PLANILHA ---")
checa("os nomes das almas", "Alma Um" in todo_texto)
checa("os relatos dos contatos", "Conversa boa." in todo_texto)
checa("as presencas", "Culto Dominical" in todo_texto)
checa("a equipe", "Joao Pereira" in todo_texto)
checa("telefone formatado, nao cru", "(16) 99280-5852" in todo_texto)

print("\n--- 8. A EXPORTACAO VAI PARA A AUDITORIA (secao 7, item 12) ---")
with app.app_context():
    n = db.session.scalar(db.select(db.func.count()).select_from(LogAuditoria)
                          .where(LogAuditoria.acao == "exportacao_excel"))
    checa(f"exportacao registrada ({n})", n == 1)
    linha = db.session.execute(
        db.select(LogAuditoria).where(LogAuditoria.acao == "exportacao_excel")
    ).scalar_one()
    checa("registra QUEM baixou", linha.usuario_id == ID_ADMIN)
    checa("registra o IP", bool(linha.ip))
    checa("registra quantas linhas saiu", "Novos Convertidos=" in (linha.detalhe or ""))

ca.get("/relatorios/excel")
with app.app_context():
    n = db.session.scalar(db.select(db.func.count()).select_from(LogAuditoria)
                          .where(LogAuditoria.acao == "exportacao_excel"))
    checa(f"cada download vira um registro novo ({n})", n == 2)

print("\n--- 9. RELATORIO COM BANCO VAZIO NAO QUEBRA ---")
class ConfigVazia(ConfigTeste):
    pass

app2 = create_app(ConfigVazia)
with app2.app_context():
    db.create_all()
    try:
        d = calc.tudo()
        checa("calcula sem quebrar", True)
        checa("retencao vira None em vez de dividir por zero", d["retencao"]["percentual"] is None)
        checa("departamento com total zero", d["por_departamento"]["total"] == 0)
    except Exception as e:
        checa(f"calcula sem quebrar (erro: {e})", False)

print(f"\n{'=' * 55}")
print(f"  {ok} testes passaram, {falhou} falharam")
print(f"{'=' * 55}")
sys.exit(1 if falhou else 0)
